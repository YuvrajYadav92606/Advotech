import io
from typing import List, Tuple
import docx
from app.schemas.document import DocumentSegment

def extract_text_from_docx(file_bytes: bytes) -> Tuple[List[DocumentSegment], int, bool]:
    """
    Extracts structured paragraphs and tables from DOCX bytes using python-docx.
    Preserves paragraph and table references.
    
    Returns:
        segments: List of DocumentSegment (paragraphs and tables)
        total_chars: Total count of extracted characters
        has_readable_text: True if readable text is found
    """
    segments: List[DocumentSegment] = []
    total_chars = 0
    segment_idx = 0
    
    doc = docx.Document(io.BytesIO(file_bytes))
    
    # 1. Extract non-empty paragraphs
    para_num = 1
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            char_count = len(text)
            total_chars += char_count
            segments.append(
                DocumentSegment(
                    segment_index=segment_idx,
                    reference_type="paragraph",
                    reference_label=f"Paragraph {para_num}",
                    content=text,
                    char_count=char_count
                )
            )
            segment_idx += 1
            para_num += 1
            
    # 2. Extract tables (tabular legal clauses, schedules, claims)
    table_num = 1
    for table in doc.tables:
        rows_text = []
        for row in table.rows:
            cell_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cell_texts:
                # Deduplicate merged cells text in same row
                seen = []
                for ct in cell_texts:
                    if not seen or seen[-1] != ct:
                        seen.append(ct)
                rows_text.append(" | ".join(seen))
                
        if rows_text:
            table_content = "\n".join(rows_text)
            char_count = len(table_content)
            total_chars += char_count
            segments.append(
                DocumentSegment(
                    segment_index=segment_idx,
                    reference_type="table",
                    reference_label=f"Table {table_num}",
                    content=table_content,
                    char_count=char_count
                )
            )
            segment_idx += 1
            table_num += 1
            
    has_readable_text = total_chars > 20
    return segments, total_chars, has_readable_text
