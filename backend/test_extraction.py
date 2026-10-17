import io
import sys
import pymupdf
import docx
from fastapi.testclient import TestClient

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

from app.main import app

def test_backend_end_to_end():
    print("=== Testing Advo टेक Backend Health & Document Processing ===")
    client = TestClient(app)
    
    # 1. Test Health Endpoint
    response = client.get("/api/v1/health")
    assert response.status_code == 200, f"Health check failed: {response.text}"
    health_data = response.json()
    print(f"[PASS] Health check: {health_data['project']} ({health_data['status']})")
    assert health_data["project"] == "Advo टेक"

    # 2. Test Root Endpoint
    response = client.get("/")
    assert response.status_code == 200
    print("[PASS] Root endpoint accessible")

    # 3. Create a synthetic Legal Notice PDF using PyMuPDF
    pdf_doc = pymupdf.open()
    
    # Page 1
    page1 = pdf_doc.new_page()
    text_page1 = (
        "LEGAL NOTICE UNDER SECTION 138 OF THE NEGOTIABLE INSTRUMENTS ACT, 1881\n\n"
        "To: Mr. Ramesh Kumar, Resident of Sector 14, Gurugram, Haryana.\n"
        "From: Adv. Suresh Sharma, High Court of Delhi, on behalf of M/s Apex Enterprises.\n\n"
        "Sir, under instructions from my client, I hereby issue this legal notice regarding dishonour of Cheque No. 452109 "
        "dated 15th January 2024 for an amount of Rs. 2,50,000/- drawn on State Bank of India.\n"
        "You are hereby called upon to remit the said sum of Rs. 2,50,000/- within 15 days of receipt of this notice."
    )
    page1.insert_text((50, 72), text_page1, fontsize=11)
    
    # Page 2
    page2 = pdf_doc.new_page()
    text_page2 = (
        "Failure to make payment within the stipulated 15 days period will compel my client to initiate criminal proceedings "
        "under Section 138 of the Negotiable Instruments Act, 1881 and civil recovery in appropriate Court of Law.\n\n"
        "Dated: 20th January 2024\nPlace: New Delhi\n(Adv. Suresh Sharma)"
    )
    page2.insert_text((50, 72), text_page2, fontsize=11)
    
    pdf_bytes = pdf_doc.write()
    pdf_doc.close()

    # Test PDF Upload Endpoint
    pdf_upload_response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("sample_legal_notice.pdf", pdf_bytes, "application/pdf")}
    )
    assert pdf_upload_response.status_code == 200, f"PDF upload failed: {pdf_upload_response.text}"
    pdf_data = pdf_upload_response.json()
    print(f"[PASS] PDF Extraction: {pdf_data['filename']}")
    print(f"       Total segments: {pdf_data['total_segments']} pages")
    print(f"       Total chars: {pdf_data['total_characters']}")
    print(f"       Has readable text: {pdf_data['has_readable_text']}")
    print(f"       Segment 1 label: {pdf_data['segments'][0]['reference_label']}")
    print(f"       Segment 2 label: {pdf_data['segments'][1]['reference_label']}")
    assert pdf_data["total_segments"] == 2
    assert pdf_data["has_readable_text"] is True

    # 4. Create a synthetic Commercial Agreement DOCX using python-docx
    docx_doc = docx.Document()
    docx_doc.add_heading("NON-DISCLOSURE AND SERVICE AGREEMENT", level=1)
    docx_doc.add_paragraph("This Agreement is entered into on 1st March 2024 between TechNova Solutions Pvt Ltd and Zenith Analytics LLP.")
    docx_doc.add_paragraph("Clause 4: Confidentiality Obligation. The Recipient shall maintain strict confidentiality for a period of 3 years.")
    
    # Add a table
    table = docx_doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Milestone'
    hdr_cells[1].text = 'Due Date'
    hdr_cells[2].text = 'Payment Amount'
    row_cells = table.add_row().cells
    row_cells[0].text = 'Phase 1 Delivery'
    row_cells[1].text = '31st March 2024'
    row_cells[2].text = 'INR 5,00,000'

    docx_stream = io.BytesIO()
    docx_doc.save(docx_stream)
    docx_bytes = docx_stream.getvalue()

    # Test DOCX Upload Endpoint
    docx_upload_response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("sample_contract.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert docx_upload_response.status_code == 200, f"DOCX upload failed: {docx_upload_response.text}"
    docx_data = docx_upload_response.json()
    print(f"[PASS] DOCX Extraction: {docx_data['filename']}")
    print(f"       Total segments: {docx_data['total_segments']}")
    print(f"       Total chars: {docx_data['total_characters']}")
    print(f"       Has readable text: {docx_data['has_readable_text']}")
    print(f"       Preview: {docx_data['preview'][:100]}...")
    assert docx_data["has_readable_text"] is True

    # 5. Test Invalid File Extension Error Handling
    invalid_response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("malicious.exe", b"executable_code", "application/octet-stream")}
    )
    assert invalid_response.status_code == 415
    print("[PASS] 415 Unsupported Media Type for invalid extensions validated")

    # 6. Test Empty File Error Handling
    empty_response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("empty.pdf", b"", "application/pdf")}
    )
    assert empty_response.status_code == 400
    print("[PASS] 400 Bad Request for empty files validated")

    print("\n>>> ALL TESTS PASSED SUCCESSFULLY! MILESTONE 1 VERIFIED. <<<")

if __name__ == "__main__":
    test_backend_end_to_end()
